import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

md_path = r"C:\Users\sanja\.gemini\antigravity-ide\brain\38dff86f-16ba-424a-b915-92c194a09fd5\research_paper_documentation.md"
txt_path = r"c:\VoiceVision-voice\voice_age_prediction\research_paper_documentation.txt"
docx_path = r"c:\VoiceVision-voice\voice_age_prediction\research_paper_documentation.docx"

with open(md_path, "r", encoding="utf-8") as f:
    md_content = f.read()

# --- 1. Save TXT File ---
with open(txt_path, "w", encoding="utf-8") as f:
    f.write(md_content)
print(f"Saved plain text documentation to: {txt_path}")

# --- 2. Create DOCX File ---
doc = docx.Document()

# Set standard margins (1 inch)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Helper function for setting cell background color
def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

lines = md_content.splitlines()

in_table = False
table_rows = []

i = 0
while i < len(lines):
    line = lines[i].strip()
    
    if not line:
        i += 1
        continue
    
    # Check table row
    if line.startswith("|") and line.endswith("|"):
        # Check if line is header separator like |---|---|
        if "---" in line:
            i += 1
            continue
        cells = [c.strip() for c in line.split("|")[1:-1]]
        if not in_table:
            in_table = True
            table_rows = []
        table_rows.append(cells)
        i += 1
        continue
    else:
        if in_table:
            # End of table, render table
            if table_rows:
                num_cols = max(len(r) for r in table_rows)
                t = doc.add_table(rows=len(table_rows), cols=num_cols)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row_data in enumerate(table_rows):
                    row = t.rows[r_idx]
                    is_header = (r_idx == 0)
                    for c_idx, cell_value in enumerate(row_data):
                        if c_idx < len(row.cells):
                            cell = row.cells[c_idx]
                            cell.text = cell_value
                            p = cell.paragraphs[0]
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.LEFT
                            for run in p.runs:
                                run.font.name = "Arial"
                                run.font.size = Pt(10)
                                if is_header:
                                    run.font.bold = True
                                    run.font.color.rgb = RGBColor(255, 255, 255)
                            if is_header:
                                set_cell_background(cell, "1F497D") # Navy blue
                            else:
                                if r_idx % 2 == 1:
                                    set_cell_background(cell, "F2F5F8") # Light zebra fill
            in_table = False
            table_rows = []

    # Headings
    if line.startswith("# "):
        p = doc.add_heading(line[2:], level=0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(31, 73, 125)
    elif line.startswith("## "):
        p = doc.add_heading(line[3:], level=1)
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(31, 73, 125)
    elif line.startswith("### "):
        p = doc.add_heading(line[4:], level=2)
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(59, 89, 152)
    elif line.startswith("- ") or line.startswith("* "):
        p = doc.add_paragraph(style='List Bullet')
        text = line[2:]
        # Parse inline bold **text**
        parts = re.split(r'(\*\*.*?\*\*)', text)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            else:
                r = p.add_run(part)
            r.font.name = "Arial"
            r.font.size = Pt(11)
    elif line.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        text = line[2:]
        parts = re.split(r'(\*\*.*?\*\*)', text)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            else:
                r = p.add_run(part)
            r.italic = True
            r.font.name = "Arial"
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(80, 80, 80)
    elif line.startswith("```"):
        # Code block marker
        i += 1
        code_lines = []
        while i < len(lines) and not lines[i].strip().startswith("```"):
            code_lines.append(lines[i])
            i += 1
        code_text = "\n".join(code_lines)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.right_indent = Inches(0.3)
        r = p.add_run(code_text)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(40, 40, 40)
    else:
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*)', line)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            else:
                r = p.add_run(part)
            r.font.name = "Arial"
            r.font.size = Pt(11)
            
    i += 1

# Render remaining table if any
if in_table and table_rows:
    num_cols = max(len(r) for r in table_rows)
    t = doc.add_table(rows=len(table_rows), cols=num_cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row_data in enumerate(table_rows):
        row = t.rows[r_idx]
        is_header = (r_idx == 0)
        for c_idx, cell_value in enumerate(row_data):
            if c_idx < len(row.cells):
                cell = row.cells[c_idx]
                cell.text = cell_value
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10)
                    if is_header:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                if is_header:
                    set_cell_background(cell, "1F497D")
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "F2F5F8")

try:
    doc.save(docx_path)
    print(f"Saved styled Word Document to: {docx_path}")
except PermissionError:
    alt_docx_path = r"c:\VoiceVision-voice\voice_age_prediction\research_paper_documentation_updated.docx"
    doc.save(alt_docx_path)
    print(f"File was locked by Word. Saved updated Word Document to: {alt_docx_path}")

