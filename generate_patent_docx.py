import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

md_path = r"c:\VoiceVision-voice\VoiceVision_Patent_IDF_Form_B.md"
docx_path = r"c:\VoiceVision-voice\VoiceVision_Patent_IDF_Form_B.docx"

with open(md_path, "r", encoding="utf-8") as f:
    md_content = f.read()

doc = docx.Document()

# Set standard margins (1 inch)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def add_formatted_text_to_p(p, text, default_size=Pt(11), default_color=None):
    # Parse bold and italic and inline code
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) >= 4:
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.font.name = "Consolas"
        else:
            run = p.add_run(token)
        
        run.font.name = "Calibri" if not run.font.name else run.font.name
        run.font.size = default_size
        if default_color:
            run.font.color.rgb = default_color

lines = md_content.splitlines()

in_table = False
table_rows = []
in_code = False
code_lines = []

i = 0
while i < len(lines):
    line = lines[i].rstrip()
    
    # Code block detection
    if line.startswith("```"):
        if in_code:
            in_code = False
            # Render code block
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run("\n".join(code_lines))
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(40, 40, 40)
            code_lines = []
        else:
            in_code = True
            code_lines = []
        i += 1
        continue
    
    if in_code:
        code_lines.append(line)
        i += 1
        continue

    # Table detection
    if line.startswith("|") and line.endswith("|"):
        if "---" in line:
            i += 1
            continue
        cells = [c.strip() for c in line.split("|")[1:-1]]
        if not in_table:
            in_table = True
            table_rows = []
        table_rows.append(cells)
        i += 1
        continue
    else:
        if in_table:
            # End of table, render table
            if table_rows:
                num_cols = max(len(r) for r in table_rows)
                t = doc.add_table(rows=len(table_rows), cols=num_cols)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row_data in enumerate(table_rows):
                    row = t.rows[r_idx]
                    is_header = (r_idx == 0)
                    for c_idx, cell_value in enumerate(row_data):
                        if c_idx < len(row.cells):
                            cell = row.cells[c_idx]
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.LEFT
                            add_formatted_text_to_p(p, cell_value, default_size=Pt(9.5), default_color=RGBColor(255, 255, 255) if is_header else None)
                            if is_header:
                                set_cell_background(cell, "1F4E79") # Deep Blue
                            else:
                                if r_idx % 2 == 1:
                                    set_cell_background(cell, "F2F4F7") # Light subtle shading
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            in_table = False
            table_rows = []

    if not line.strip():
        i += 1
        continue

    if line.startswith("---"):
        i += 1
        continue

    # Headings
    if line.startswith("# "):
        p = doc.add_heading(level=0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_formatted_text_to_p(p, line[2:], default_size=Pt(20), default_color=RGBColor(31, 78, 121))
    elif line.startswith("## "):
        p = doc.add_heading(level=1)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_formatted_text_to_p(p, line[3:], default_size=Pt(15), default_color=RGBColor(31, 78, 121))
    elif line.startswith("### "):
        p = doc.add_heading(level=2)
        add_formatted_text_to_p(p, line[4:], default_size=Pt(13), default_color=RGBColor(44, 90, 160))
    elif line.startswith("#### "):
        p = doc.add_heading(level=3)
        add_formatted_text_to_p(p, line[5:], default_size=Pt(11.5), default_color=RGBColor(30, 30, 30))
    elif line.startswith("##### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        add_formatted_text_to_p(p, f"**{line[6:]}**", default_size=Pt(11))
    elif line.startswith("* ") or line.startswith("- "):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        add_formatted_text_to_p(p, line[2:])
    elif re.match(r'^\d+\.\s', line):
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        text = re.sub(r'^\d+\.\s', '', line)
        add_formatted_text_to_p(p, text)
    elif line.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        add_formatted_text_to_p(p, line[2:], default_size=Pt(10.5))
    else:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        add_formatted_text_to_p(p, line)

    i += 1

# If table was open at the very end
if in_table and table_rows:
    num_cols = max(len(r) for r in table_rows)
    t = doc.add_table(rows=len(table_rows), cols=num_cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row_data in enumerate(table_rows):
        row = t.rows[r_idx]
        is_header = (r_idx == 0)
        for c_idx, cell_value in enumerate(row_data):
            if c_idx < len(row.cells):
                cell = row.cells[c_idx]
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.LEFT
                add_formatted_text_to_p(p, cell_value, default_size=Pt(9.5), default_color=RGBColor(255, 255, 255) if is_header else None)
                if is_header:
                    set_cell_background(cell, "1F4E79")
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "F2F4F7")

doc.save(docx_path)
print(f"Successfully generated DOCX at: {docx_path}")
